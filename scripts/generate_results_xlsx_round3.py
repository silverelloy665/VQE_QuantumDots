import json
from pathlib import Path
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd

import sys
REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))
XLSX_PATH = REPO_ROOT / "results.xlsx"
DATA_DIR = REPO_ROOT / "data"

SPEC_V3_PATH = DATA_DIR / "hardware_job_spec_v3.json"
CACHE_HARD_PATH = DATA_DIR / "benchmark_results_hard.json"
SUMMARY_CSV = DATA_DIR / "hard_benchmark_summary.csv"


def style_worksheet(ws, title: str):
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.0000"

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def update_results_xlsx_round3():
    print("=" * 80)
    print("PHASE J: UPDATING results.xlsx WITH ROUND 3 SHEETS")
    print("=" * 80)

    if not XLSX_PATH.exists():
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    else:
        wb = openpyxl.load_workbook(XLSX_PATH)

    # 1. Review Verification Sheet
    verif_data = [
        {"Finding #": 1, "Review Finding": "DexcG/PCU2 zero-init reaches CASCI error 0.015 mHa; summary CSV only aggregated random init", "Result": "VERIFIED & REPRODUCED (DexcG reaches 0.0149 mHa, PCU2 reaches 0.0148 mHa)", "Status": "RESOLVED"},
        {"Finding #": 2, "Review Finding": "converged flag was hardcoded True; maxiter=25 ended in unconverged state from random", "Result": "VERIFIED & FIXED (Schema v3 defines converged by optimizer success or energy plateau < 1e-8)", "Status": "RESOLVED"},
        {"Finding #": 3, "Review Finding": "Mean % correlation distorted by outliers (-391.7%)", "Result": "VERIFIED & REPLACED (All reports now use median % correlation recovered)", "Status": "RESOLVED"},
        {"Finding #": 4, "Review Finding": "run_live_hardware_job.py hardcoded params_1up (|theta| <= 3e-4 repeating)", "Result": "VERIFIED & ELIMINATED (Params loaded dynamically from converged Schema v3 cache records)", "Status": "RESOLVED"},
        {"Finding #": 5, "Review Finding": "Fidelity discrepancy (0.0398 device mean for 2-CZ vs 0.00186 for 489-CZ; ignores 1Q, RO, T1)", "Result": "VERIFIED & UNIFIED (Single error model: actual transpiled CZ edges * readout * exp(-t/T1))", "Status": "RESOLVED"},
        {"Finding #": 6, "Review Finding": "Rotosolve text claimed omega=4 with points 0, pi/4, pi/2 (singular model)", "Result": "VERIFIED & CORRECTED (FFT peak at omega=2, formula E=a+b*cos(2th)+c*sin(2th))", "Status": "RESOLVED"},
        {"Finding #": 7, "Review Finding": "figures/circuit_*.png dated Sep 21 showed old 2e2o circuits; silent GenericBackendV2 fallback", "Result": "VERIFIED & RESOLVED (2e2o preserved as _2e2o_4q.png, 8q rendered; mock fallback requires --mock)", "Status": "RESOLVED"},
    ]
    df_verif = pd.DataFrame(verif_data)
    if "Round 3 - Review Verification" in wb.sheetnames:
        del wb["Round 3 - Review Verification"]
    ws_v = wb.create_sheet("Round 3 - Review Verification")
    for r in dataframe_to_rows(df_verif, index=False, header=True):
        ws_v.append(r)
    style_worksheet(ws_v, "Review Verification")

    # Load cache records
    with open(CACHE_HARD_PATH, "r", encoding="utf-8") as f:
        cache_records = json.load(f).get("records", [])

    from src.molecule import load_hard_hamiltonian
    _, exact_energy, meta = load_hard_hamiltonian()
    ref_det_energy = float(meta["reference_determinant_energy"])
    delta_corr = abs(ref_det_energy - exact_energy)

    processed = []
    for r in cache_records:
        f_e = float(r["final_energy"])
        err_mha = abs(f_e - exact_energy) * 1000.0
        pct_corr = ((ref_det_energy - f_e) / delta_corr) * 100.0 if delta_corr > 1e-12 else 0.0
        better_than_hf = (f_e < ref_det_energy - 1e-9)
        succ = (err_mha < 1.6 and better_than_hf)
        init_raw = r["init"]
        init_group = "perturbed" if init_raw.startswith("perturbed") else init_raw
        processed.append({
            "ansatz": r["ansatz"],
            "optimizer": r["optimizer"],
            "init": init_raw,
            "init_group": init_group,
            "seed": r.get("seed", 0),
            "final_energy": f_e,
            "error_mHa": err_mha,
            "pct_corr_recovered": pct_corr,
            "better_than_HF": better_than_hf,
            "success": succ,
            "converged": bool(r.get("converged", False)),
            "hit_maxiter": bool(r.get("hit_maxiter", False)),
            "total_evaluations": int(r.get("total_evaluations", 0)),
            "iterations_used": int(r.get("iterations_used", 0)),
            "wall_time": float(r.get("wall_time", 0.0))
        })
    df_all = pd.DataFrame(processed)

    # 2. Zero Init Sheet
    sub_zero = df_all[df_all["init"] == "zero"].sort_values(by="error_mHa")
    df_zero = sub_zero[["ansatz", "optimizer", "error_mHa", "pct_corr_recovered", "total_evaluations", "converged", "wall_time"]].copy()
    df_zero.columns = ["Ansatz", "Optimizer", "Error (mHa)", "% Corr Recovered", "Evaluations", "Converged", "Wall Time (s)"]
    if "Round 3 - Zero Init" in wb.sheetnames:
        del wb["Round 3 - Zero Init"]
    ws_z = wb.create_sheet("Round 3 - Zero Init")
    for r in dataframe_to_rows(df_zero, index=False, header=True):
        ws_z.append(r)
    style_worksheet(ws_z, "Zero Init")

    # 3. Perturbed Init Sheet
    sub_pert = df_all[df_all["init_group"] == "perturbed"]
    pert_rows = []
    for (a_name, opt_name), group in sub_pert.groupby(["ansatz", "optimizer"]):
        errs = group["error_mHa"].values
        q25, med, q75 = np.percentile(errs, [25, 50, 75])
        pert_rows.append({
            "Ansatz": a_name,
            "Optimizer": opt_name,
            "N Runs": len(errs),
            "Median Error (mHa)": float(med),
            "IQR (mHa)": float(q75 - q25),
            "Min Error (mHa)": float(np.min(errs)),
            "Max Error (mHa)": float(np.max(errs)),
            "Median % Corr": float(np.median(group["pct_corr_recovered"].values)),
            "Success Rate (%)": float(np.mean(group["success"].values) * 100.0),
            "Median Evaluations": int(np.median(group["total_evaluations"].values))
        })
    df_pert = pd.DataFrame(pert_rows).sort_values(by="Median Error (mHa)")
    if "Round 3 - Perturbed Init" in wb.sheetnames:
        del wb["Round 3 - Perturbed Init"]
    ws_p = wb.create_sheet("Round 3 - Perturbed Init")
    for r in dataframe_to_rows(df_pert, index=False, header=True):
        ws_p.append(r)
    style_worksheet(ws_p, "Perturbed Init")

    # 4. Random Init Sheet
    sub_rand = df_all[df_all["init"] == "random"]
    rand_rows = []
    for (a_name, opt_name), group in sub_rand.groupby(["ansatz", "optimizer"]):
        errs = group["error_mHa"].values
        q25, med, q75 = np.percentile(errs, [25, 50, 75])
        rand_rows.append({
            "Ansatz": a_name,
            "Optimizer": opt_name,
            "N Runs": len(errs),
            "Median Error (mHa)": float(med),
            "IQR (mHa)": float(q75 - q25),
            "Min Error (mHa)": float(np.min(errs)),
            "Max Error (mHa)": float(np.max(errs)),
            "Median % Corr": float(np.median(group["pct_corr_recovered"].values)),
            "Success Rate (%)": float(np.mean(group["success"].values) * 100.0),
            "Median Evaluations": int(np.median(group["total_evaluations"].values))
        })
    df_rand = pd.DataFrame(rand_rows).sort_values(by="Median Error (mHa)")
    if "Round 3 - Random Init" in wb.sheetnames:
        del wb["Round 3 - Random Init"]
    ws_r = wb.create_sheet("Round 3 - Random Init")
    for r in dataframe_to_rows(df_rand, index=False, header=True):
        ws_r.append(r)
    style_worksheet(ws_r, "Random Init")

    # 5. Hardware Job Spec V3 Sheet
    if SPEC_V3_PATH.exists():
        with open(SPEC_V3_PATH, "r", encoding="utf-8") as f:
            v3_specs = json.load(f)
        hw_rows = []
        for s in v3_specs:
            hw_rows.append({
                "Priority": s["priority"],
                "Experiment Name": s["experiment_name"],
                "Resilience Type": s["resilience_type"],
                "Backend": s["backend"],
                "Source": s.get("backend_source", "LIVE"),
                "Calibration Time": s.get("calibration_timestamp", "ASSUMED"),
                "QWC Circuits": s["qwc_circuits"],
                "2Q Gates": s["transpiled_2q_gates"],
                "Depth": s["transpiled_depth"],
                "Estimated Fidelity": round(s["estimated_fidelity"], 4),
                "Chemistry Signal": s.get("has_chemistry_signal", False),
                "Recommendation": s["recommendation"]
            })
        df_hw = pd.DataFrame(hw_rows)
        if "Round 3 - Hardware Jobs V3" in wb.sheetnames:
            del wb["Round 3 - Hardware Jobs V3"]
        ws_h = wb.create_sheet("Round 3 - Hardware Jobs V3")
        for r in dataframe_to_rows(df_hw, index=False, header=True):
            ws_h.append(r)
        style_worksheet(ws_h, "Hardware Jobs V3")

    wb.save(XLSX_PATH)
    print(f"Successfully saved updated results.xlsx with Round 3 sheets: {[s for s in wb.sheetnames if 'Round 3' in s]}")


if __name__ == "__main__":
    update_results_xlsx_round3()
